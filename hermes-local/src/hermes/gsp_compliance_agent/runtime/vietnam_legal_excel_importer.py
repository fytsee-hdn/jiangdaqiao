"""
vietnam_legal_excel_importer.py — GSP Compliance Agent: Vietnam legal Excel import.

Reads a Vietnam HSE/ESG legal index Excel workbook, classifies sheets,
extracts legal_source_register and legal_curated_index records,
and writes them to JSONL under the compliance data root.

Key principles:
- dry_run=True by default.
- execute=False by default.
- All records remain pending review / not_approved.
- No legal judgment, no applicability analysis, no GSP modification.
"""

from __future__ import annotations

import json
import os
import re
from datetime import datetime, timezone
from hashlib import sha256
from typing import Any, Dict, List, Optional, Tuple

from hermes.gsp_compliance_agent.runtime.compliance_data_paths import (
    resolve_compliance_path,
    ensure_empty_jsonl,
)


# ══════════════════════════════════════════════════════════════════
#  Constants
# ══════════════════════════════════════════════════════════════════

# Sheets that map to legal_source_register
SOURCE_REGISTER_SHEETS = {"VBQPPL", "TCVN"}

# Sheets that map to legal_curated_index
CURATED_INDEX_SHEETS = {
    "ESG-E", "ESG-S", "ESG-G",
    "CV YÊU CẦU NGHIÊM NGẶT",
    # Topic sheets with numeric prefix pattern
}

SHEET_FUNCTION_MAP = {
    "ESG-E": "environment",
    "ESG-S": "labour",
    "ESG-G": "management",
    "CV YÊU CẦU NGHIÊM NGẶT": "safety",
}

SHEET_RISK_MAP = {
    "ESG-E": "environmental_compliance",
    "ESG-S": "occupational_safety",
    "ESG-G": "corporate_governance",
    "CV YÊU CẦU NGHIÊM NGẶT": "occupational_safety",
}

SOURCE_TYPE_ALIASES = {
    "law": "law",
    "luật": "law",
    "decree": "decree",
    "nghị định": "decree",
    "circular": "circular",
    "thông tư": "circular",
    "decision": "other",
    "quyết định": "other",
}

STATUS_ALIASES = {
    "còn hiệu lực": "current",
    "đã sửa đổi": "amended",
    "hết hiệu lực": "repealed",
    "đã thay thế": "replaced",
    "current": "current",
    "amended": "amended",
    "repealed": "repealed",
}

COUNTRY_CODE = "VN"
JURISDICTION = "national"
LANGUAGE = "vi"
SOURCE_ORIGIN = "user_curated_legal_index_excel"


# ══════════════════════════════════════════════════════════════════
#  Paths
# ══════════════════════════════════════════════════════════════════


def legal_source_register_path() -> str:
    return resolve_compliance_path("legal/legal_source_register/legal_source_register.jsonl")


def legal_curated_index_path() -> str:
    return resolve_compliance_path("legal/legal_curated_index/legal_curated_index.jsonl")


def country_profiles_path() -> str:
    return resolve_compliance_path("legal/country_profiles/country_legal_profile.jsonl")


def reports_dir() -> str:
    return resolve_compliance_path("legal/reports")


# ══════════════════════════════════════════════════════════════════
#  Helpers
# ══════════════════════════════════════════════════════════════════


def _load_jsonl(path: str) -> List[Dict[str, Any]]:
    """Load JSONL records from path."""
    records = []
    if not os.path.isfile(path):
        return records
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    return records


def _cell(val: Any) -> str:
    """Convert cell value to stripped string."""
    if val is None:
        return ""
    return str(val).strip()


def _normalize_source_type(raw: str) -> str:
    """Normalize Vietnamese/English source type to schema enum."""
    key = raw.strip().lower()
    if key in SOURCE_TYPE_ALIASES:
        return SOURCE_TYPE_ALIASES[key]
    # Fall through: try direct match
    if key in ("law", "decree", "circular", "technical_regulation", "standard",
               "official_guidance", "legal_index", "curated_excel_index", "other"):
        return key
    return "other"


def _normalize_status(raw: str) -> str:
    """Normalize Vietnamese status to schema enum."""
    key = raw.strip().lower()
    if key in STATUS_ALIASES:
        return STATUS_ALIASES[key]
    return "current"


def _is_topic_sheet(name: str) -> bool:
    """Check if sheet name matches topic sheet pattern like '01. TỔ CHỨC'."""
    return bool(re.match(r"^\d{2}\.\s", name.strip()))


def _infer_function_from_topic_sheet(name: str) -> str:
    """Infer function from topic sheet name."""
    upper = name.upper()
    if "TỔ CHỨC" in upper:
        return "management"
    if "KẾ HOẠCH" in upper:
        return "management"
    if "QLRR" in upper or "RỦI RO" in upper:
        return "safety"
    return "other"


def _generate_id(prefix: str, counter: int) -> str:
    """Generate a unique record ID."""
    return f"{prefix}-{COUNTRY_CODE}-{counter:04d}"


# ══════════════════════════════════════════════════════════════════
#  Workbook Inspection
# ══════════════════════════════════════════════════════════════════


def inspect_vietnam_legal_workbook(file_path: str) -> Dict[str, Any]:
    """Inspect a Vietnam legal Excel workbook and return structure summary.

    Args:
        file_path: Path to the .xlsx file.

    Returns:
        Dict with keys: filename, sheets (list), total_sheets, warnings.
        Each sheet entry: name, max_row, max_col, headers (list), row_count, classification.
    """
    try:
        import openpyxl
    except ImportError:
        return {
            "filename": os.path.basename(file_path),
            "error": "openpyxl not installed",
            "sheets": [], "total_sheets": 0,
        }

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets = []
    warnings = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row = 0
        headers = []
        try:
            # Read headers from first row
            for cell in ws[1]:
                headers.append(str(cell.value).strip() if cell.value is not None else "")
            # Count non-empty rows
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row or 0, 500)):
                if any(cell.value is not None for cell in row):
                    max_row += 1
        except Exception:
            pass

        classification = classify_sheet(sheet_name, headers)

        sheets.append({
            "name": sheet_name,
            "max_row": max_row + 1 if max_row else 0,
            "max_col": len(headers),
            "row_count": max_row,
            "headers": headers,
            "classification": classification,
        })

    wb.close()

    return {
        "filename": os.path.basename(file_path),
        "file_path": file_path,
        "total_sheets": len(sheets),
        "sheets": sheets,
        "warnings": warnings,
    }


# ══════════════════════════════════════════════════════════════════
#  Sheet Classification
# ══════════════════════════════════════════════════════════════════


def classify_sheet(sheet_name: str, headers: List[str]) -> str:
    """Classify a sheet into the target import object.

    Returns one of: 'legal_source_register', 'legal_curated_index', 'unknown'.
    """
    name_upper = sheet_name.strip().upper()

    if name_upper in {s.upper() for s in SOURCE_REGISTER_SHEETS}:
        return "legal_source_register"

    if name_upper in {s.upper() for s in CURATED_INDEX_SHEETS}:
        return "legal_curated_index"

    if _is_topic_sheet(sheet_name):
        return "legal_curated_index"

    # Fallback: classify by header content
    header_text = " ".join(h.lower() for h in headers if h)
    if "luật" in header_text or "văn bản" in header_text or "số hiệu" in header_text:
        if "yêu cầu" in header_text or "mô tả" in header_text:
            return "legal_curated_index"
        return "legal_source_register"

    return "unknown"


# ══════════════════════════════════════════════════════════════════
#  Parse VBQPPL Sheet
# ══════════════════════════════════════════════════════════════════


def _parse_vbqppl_sheet(ws, sheet_name: str, counter: int) -> Tuple[List[Dict], int, List[str]]:
    """Parse VBQPPL sheet into legal_source_register records.

    Returns: (records, next_counter, warnings)
    """
    records = []
    warnings = []
    header_row = 1

    # Find header mapping
    headers = {}
    for col in range(1, ws.max_column + 1):
        val = _cell(ws.cell(row=header_row, column=col).value)
        headers[col] = val

    # Column detection
    col_map = {}
    for col, h in headers.items():
        hu = h.upper()
        if hu == "SỐ HIỆU VĂN BẢN":
            col_map["law_number"] = col
        elif hu == "LOẠI VĂN BẢN":
            col_map["source_type"] = col
        elif hu == "TÊN VĂN BẢN":
            col_map["law_name_original"] = col
        elif hu == "TÊN TIẾNG ANH":
            col_map["law_name_en"] = col
        elif hu == "CƠ QUAN BAN HÀNH":
            col_map["issuing_authority"] = col
        elif hu == "NGÀY BAN HÀNH":
            col_map["issue_date"] = col
        elif hu == "NGÀY HIỆU LỰC":
            col_map["effective_date"] = col
        elif hu == "HIỆU LỰC":
            col_map["status"] = col
        elif hu == "THUỘC LĨNH VỰC":
            col_map["category"] = col
        elif hu == "GHI CHÚ":
            col_map["notes"] = col

    for row in range(header_row + 1, ws.max_row + 1):
        law_number = _cell(ws.cell(row=row, column=col_map.get("law_number", 0)).value) if "law_number" in col_map else ""
        if not law_number:
            continue

        source_type_raw = _cell(ws.cell(row=row, column=col_map.get("source_type", 0)).value) if "source_type" in col_map else ""

        record = {
            "legal_source_id": _generate_id("LS", counter),
            "country_code": COUNTRY_CODE,
            "jurisdiction": JURISDICTION,
            "source_type": _normalize_source_type(source_type_raw),
            "law_number": law_number,
            "law_name_original": _cell(ws.cell(row=row, column=col_map.get("law_name_original", 0)).value) if "law_name_original" in col_map else "",
            "law_name_en": _cell(ws.cell(row=row, column=col_map.get("law_name_en", 0)).value) if "law_name_en" in col_map else "",
            "issuing_authority": _cell(ws.cell(row=row, column=col_map.get("issuing_authority", 0)).value) if "issuing_authority" in col_map else "",
            "issue_date": _cell(ws.cell(row=row, column=col_map.get("issue_date", 0)).value) if "issue_date" in col_map else "",
            "effective_date": _cell(ws.cell(row=row, column=col_map.get("effective_date", 0)).value) if "effective_date" in col_map else "",
            "status": _normalize_status(_cell(ws.cell(row=row, column=col_map.get("status", 0)).value) if "status" in col_map else ""),
            "language": LANGUAGE,
            "source_file_id": "",
            "source_file_hash": "",
            "source_url": "",
            "source_origin": SOURCE_ORIGIN,
            "replaces": "",
            "amended_by": [],
            "related_sources": [],
            "review_status": "registered",
            "approval_status": "not_approved",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        # Add notes
        notes_parts = []
        if "category" in col_map:
            cat = _cell(ws.cell(row=row, column=col_map["category"]).value)
            if cat:
                notes_parts.append(f"Category: {cat}")
        if "notes" in col_map:
            notes = _cell(ws.cell(row=row, column=col_map["notes"]).value)
            if notes:
                notes_parts.append(notes)
        if notes_parts:
            record["notes"] = "; ".join(notes_parts)

        records.append(record)
        counter += 1

    return records, counter, warnings


# ══════════════════════════════════════════════════════════════════
#  Parse TCVN Sheet
# ══════════════════════════════════════════════════════════════════


def _parse_tcvn_sheet(ws, sheet_name: str, counter: int) -> Tuple[List[Dict], int, List[str]]:
    """Parse TCVN sheet into legal_source_register records (source_type=technical_regulation)."""
    records = []
    warnings = []
    header_row = 1

    headers = {}
    for col in range(1, ws.max_column + 1):
        val = _cell(ws.cell(row=header_row, column=col).value)
        headers[col] = val

    col_map = {}
    for col, h in headers.items():
        hu = h.upper()
        if hu == "SỐ HIỆU TCVN" or "TCVN" in hu:
            col_map["law_number"] = col
        elif hu == "TÊN TIÊU CHUẨN":
            col_map["law_name_original"] = col
        elif hu == "TÊN TIẾNG ANH":
            col_map["law_name_en"] = col
        elif hu == "CƠ QUAN BAN HÀNH":
            col_map["issuing_authority"] = col
        elif hu == "NĂM BAN HÀNH":
            col_map["issue_date"] = col
        elif hu == "LĨNH VỰC":
            col_map["category"] = col
        elif hu == "GHI CHÚ":
            col_map["notes"] = col

    for row in range(header_row + 1, ws.max_row + 1):
        law_number = _cell(ws.cell(row=row, column=col_map.get("law_number", 0)).value) if "law_number" in col_map else ""
        if not law_number:
            continue

        record = {
            "legal_source_id": _generate_id("LS", counter),
            "country_code": COUNTRY_CODE,
            "jurisdiction": JURISDICTION,
            "source_type": "technical_regulation",
            "law_number": law_number,
            "law_name_original": _cell(ws.cell(row=row, column=col_map.get("law_name_original", 0)).value) if "law_name_original" in col_map else "",
            "law_name_en": _cell(ws.cell(row=row, column=col_map.get("law_name_en", 0)).value) if "law_name_en" in col_map else "",
            "issuing_authority": _cell(ws.cell(row=row, column=col_map.get("issuing_authority", 0)).value) if "issuing_authority" in col_map else "",
            "issue_date": _cell(ws.cell(row=row, column=col_map.get("issue_date", 0)).value) if "issue_date" in col_map else "",
            "effective_date": "",
            "status": "current",
            "language": LANGUAGE,
            "source_file_id": "",
            "source_file_hash": "",
            "source_url": "",
            "source_origin": SOURCE_ORIGIN,
            "replaces": "",
            "amended_by": [],
            "related_sources": [],
            "review_status": "registered",
            "approval_status": "not_approved",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        notes_parts = []
        if "category" in col_map:
            cat = _cell(ws.cell(row=row, column=col_map["category"]).value)
            if cat:
                notes_parts.append(f"Category: {cat}")
        if "notes" in col_map:
            notes = _cell(ws.cell(row=row, column=col_map["notes"]).value)
            if notes:
                notes_parts.append(notes)
        if notes_parts:
            record["notes"] = "; ".join(notes_parts)

        records.append(record)
        counter += 1

    return records, counter, warnings


# ══════════════════════════════════════════════════════════════════
#  Parse ESG / Curated Index Sheets
# ══════════════════════════════════════════════════════════════════


def _parse_esg_sheet(ws, sheet_name: str, counter: int) -> Tuple[List[Dict], int, List[str]]:
    """Parse ESG-E/ESG-S/ESG-G sheet into legal_curated_index records."""
    records = []
    warnings = []
    header_row = 1

    headers = {}
    for col in range(1, ws.max_column + 1):
        val = _cell(ws.cell(row=header_row, column=col).value)
        headers[col] = val

    col_map = {}
    for col, h in headers.items():
        hu = h.upper().replace(" ", "")
        if hu == "CHỦĐỀ":
            col_map["topic"] = col
        elif hu == "SỐHIỆUVB":
            col_map["law_number"] = col
        elif hu == "TÊNVĂNBẢN":
            col_map["law_name"] = col
        elif hu == "ĐIỀUKHOẢN":
            col_map["article_reference"] = col
        elif hu in ("YÊUCẦU/MÔTẢ", "YÊUCẦUMÔTẢ"):
            col_map["original_text"] = col
        elif hu == "ÁPDỤNGCHO":
            col_map["applicability_hint"] = col
        elif hu == "RỦIRO":
            col_map["risk_area_hint"] = col
        elif hu == "GHICHÚ":
            col_map["notes"] = col

    function_hint = SHEET_FUNCTION_MAP.get(sheet_name.upper(), "other")
    risk_area = SHEET_RISK_MAP.get(sheet_name.upper(), "")

    for row in range(header_row + 1, ws.max_row + 1):
        topic = _cell(ws.cell(row=row, column=col_map.get("topic", 0)).value) if "topic" in col_map else ""
        original_text = _cell(ws.cell(row=row, column=col_map.get("original_text", 0)).value) if "original_text" in col_map else ""

        if not topic and not original_text:
            continue  # Skip empty rows

        record = {
            "legal_index_id": _generate_id("LI", counter),
            "legal_source_id": "",
            "country_code": COUNTRY_CODE,
            "index_source_type": sheet_name.strip().upper().replace(" ", "_"),
            "uploaded_file_id": "",
            "uploaded_filename": "",
            "sheet_name": sheet_name,
            "row_number": row,
            "topic": topic or original_text[:100],
            "subtopic": sheet_name,
            "law_number": _cell(ws.cell(row=row, column=col_map.get("law_number", 0)).value) if "law_number" in col_map else "",
            "law_name": _cell(ws.cell(row=row, column=col_map.get("law_name", 0)).value) if "law_name" in col_map else "",
            "article_reference": _cell(ws.cell(row=row, column=col_map.get("article_reference", 0)).value) if "article_reference" in col_map else "",
            "original_text_or_summary": original_text,
            "applicability_hint": _cell(ws.cell(row=row, column=col_map.get("applicability_hint", 0)).value) if "applicability_hint" in col_map else "",
            "function_hint": function_hint,
            "risk_area_hint": _cell(ws.cell(row=row, column=col_map.get("risk_area_hint", 0)).value) if "risk_area_hint" in col_map else risk_area,
            "review_status": "indexed_pending_analysis",
            "approval_status": "not_approved",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        records.append(record)
        counter += 1

    return records, counter, warnings


def _parse_strict_requirement_sheet(
    ws, sheet_name: str, counter: int
) -> Tuple[List[Dict], int, List[str]]:
    """Parse 'CV YÊU CẦU NGHIÊM NGẶT' sheet."""
    records = []
    warnings = []
    header_row = 1

    headers = {}
    for col in range(1, ws.max_column + 1):
        val = _cell(ws.cell(row=header_row, column=col).value)
        headers[col] = val

    col_map = {}
    for col, h in headers.items():
        hu = h.upper().replace(" ", "")
        if hu == "HẠNGMỤC":
            col_map["topic"] = col
        elif hu == "SỐHIỆUVB":
            col_map["law_number"] = col
        elif hu == "ĐIỀUKHOẢN":
            col_map["article_reference"] = col
        elif hu == "YÊUCẦUNGHIÊMNGẶT":
            col_map["original_text"] = col
        elif hu == "TẦNSUẤTKIỂMTRA":
            col_map["applicability_hint"] = col
        elif hu == "GHICHÚ":
            col_map["notes"] = col

    for row in range(header_row + 1, ws.max_row + 1):
        topic = _cell(ws.cell(row=row, column=col_map.get("topic", 0)).value) if "topic" in col_map else ""
        original_text = _cell(ws.cell(row=row, column=col_map.get("original_text", 0)).value) if "original_text" in col_map else ""

        if not topic and not original_text:
            continue

        record = {
            "legal_index_id": _generate_id("LI", counter),
            "legal_source_id": "",
            "country_code": COUNTRY_CODE,
            "index_source_type": "CV_YEU_CAU_NGHIEM_NGAT",
            "uploaded_file_id": "",
            "uploaded_filename": "",
            "sheet_name": sheet_name,
            "row_number": row,
            "topic": topic or original_text[:100],
            "subtopic": "strict_safety_requirement",
            "law_number": _cell(ws.cell(row=row, column=col_map.get("law_number", 0)).value) if "law_number" in col_map else "",
            "law_name": "",
            "article_reference": _cell(ws.cell(row=row, column=col_map.get("article_reference", 0)).value) if "article_reference" in col_map else "",
            "original_text_or_summary": original_text,
            "applicability_hint": _cell(ws.cell(row=row, column=col_map.get("applicability_hint", 0)).value) if "applicability_hint" in col_map else "",
            "function_hint": "safety",
            "risk_area_hint": "occupational_safety",
            "review_status": "indexed_pending_analysis",
            "approval_status": "not_approved",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        records.append(record)
        counter += 1

    return records, counter, warnings


def _parse_topic_sheet(ws, sheet_name: str, counter: int) -> Tuple[List[Dict], int, List[str]]:
    """Parse topic sheets (01. TỔ CHỨC, 02. KẾ HOẠCH, etc.)."""
    records = []
    warnings = []
    header_row = 1

    headers = {}
    for col in range(1, ws.max_column + 1):
        val = _cell(ws.cell(row=header_row, column=col).value)
        headers[col] = val

    col_map = {}
    for col, h in headers.items():
        hu = h.upper()
        if hu == "YÊU CẦU":
            col_map["topic"] = col
        elif hu == "CĂN CỨ PHÁP LÝ":
            col_map["law_number"] = col
        elif hu == "ĐIỀU KHOẢN":
            col_map["article_reference"] = col
        elif hu == "MÔ TẢ":
            col_map["original_text"] = col
        elif hu == "GHI CHÚ":
            col_map["notes"] = col

    function_hint = _infer_function_from_topic_sheet(sheet_name)

    for row in range(header_row + 1, ws.max_row + 1):
        topic = _cell(ws.cell(row=row, column=col_map.get("topic", 0)).value) if "topic" in col_map else ""
        original_text = _cell(ws.cell(row=row, column=col_map.get("original_text", 0)).value) if "original_text" in col_map else ""

        if not topic and not original_text:
            continue

        record = {
            "legal_index_id": _generate_id("LI", counter),
            "legal_source_id": "",
            "country_code": COUNTRY_CODE,
            "index_source_type": "topic_sheet",
            "uploaded_file_id": "",
            "uploaded_filename": "",
            "sheet_name": sheet_name,
            "row_number": row,
            "topic": topic or original_text[:100],
            "subtopic": sheet_name.strip(),
            "law_number": _cell(ws.cell(row=row, column=col_map.get("law_number", 0)).value) if "law_number" in col_map else "",
            "law_name": "",
            "article_reference": _cell(ws.cell(row=row, column=col_map.get("article_reference", 0)).value) if "article_reference" in col_map else "",
            "original_text_or_summary": original_text,
            "applicability_hint": "",
            "function_hint": function_hint,
            "risk_area_hint": "",
            "review_status": "indexed_pending_analysis",
            "approval_status": "not_approved",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        records.append(record)
        counter += 1

    return records, counter, warnings


# ══════════════════════════════════════════════════════════════════
#  Sheet Router
# ══════════════════════════════════════════════════════════════════


def _parse_sheet(ws, sheet_name: str, counter: int) -> Tuple[str, List[Dict], int, List[str]]:
    """Route sheet to appropriate parser.

    Returns: (target_type, records, next_counter, warnings)
    """
    name_upper = sheet_name.strip().upper()
    classification = classify_sheet(sheet_name, [])

    if name_upper == "VBQPPL":
        records, counter, warnings = _parse_vbqppl_sheet(ws, sheet_name, counter)
        return "legal_source_register", records, counter, warnings
    elif name_upper == "TCVN":
        records, counter, warnings = _parse_tcvn_sheet(ws, sheet_name, counter)
        return "legal_source_register", records, counter, warnings
    elif name_upper in {"ESG-E", "ESG-S", "ESG-G"}:
        records, counter, warnings = _parse_esg_sheet(ws, sheet_name, counter)
        return "legal_curated_index", records, counter, warnings
    elif name_upper == "CV YÊU CẦU NGHIÊM NGẶT":
        records, counter, warnings = _parse_strict_requirement_sheet(ws, sheet_name, counter)
        return "legal_curated_index", records, counter, warnings
    elif _is_topic_sheet(sheet_name):
        records, counter, warnings = _parse_topic_sheet(ws, sheet_name, counter)
        return "legal_curated_index", records, counter, warnings
    else:
        return "unknown", [], counter, []


# ══════════════════════════════════════════════════════════════════
#  Main Parse Functions
# ══════════════════════════════════════════════════════════════════


def parse_legal_source_register_rows(
    file_path: str,
    sheet_mapping: Optional[Dict[str, Any]] = None,
    max_rows: Optional[int] = None,
) -> Tuple[bool, str, List[Dict[str, Any]], List[str]]:
    """Parse legal_source_register records from the workbook.

    Args:
        file_path: Path to .xlsx file.
        sheet_mapping: Optional mapping overrides.
        max_rows: Optional max row limit.

    Returns: (success, message, records, warnings)
    """
    try:
        import openpyxl
    except ImportError:
        return False, "openpyxl not installed", [], ["openpyxl required"]

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    all_records = []
    all_warnings = []
    counter = _load_jsonl(legal_source_register_path())
    counter = len(counter) + 1

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        target, records, counter, warnings = _parse_sheet(ws, sheet_name, counter)
        if target == "legal_source_register":
            all_records.extend(records)
            all_warnings.extend(warnings)
            if max_rows and len(all_records) >= max_rows:
                all_records = all_records[:max_rows]
                break

    wb.close()
    return True, f"Parsed {len(all_records)} legal_source_register records.", all_records, all_warnings


def parse_legal_curated_index_rows(
    file_path: str,
    sheet_mapping: Optional[Dict[str, Any]] = None,
    max_rows: Optional[int] = None,
) -> Tuple[bool, str, List[Dict[str, Any]], List[str]]:
    """Parse legal_curated_index records from the workbook."""
    try:
        import openpyxl
    except ImportError:
        return False, "openpyxl not installed", [], ["openpyxl required"]

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    all_records = []
    all_warnings = []
    counter = _load_jsonl(legal_curated_index_path())
    counter = len(counter) + 1

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        target, records, counter, warnings = _parse_sheet(ws, sheet_name, counter)
        if target == "legal_curated_index":
            all_records.extend(records)
            all_warnings.extend(warnings)
            if max_rows and len(all_records) >= max_rows:
                all_records = all_records[:max_rows]
                break

    wb.close()
    return True, f"Parsed {len(all_records)} legal_curated_index records.", all_records, all_warnings


# ══════════════════════════════════════════════════════════════════
#  Validation
# ══════════════════════════════════════════════════════════════════


def validate_legal_source_register_record(record: Dict[str, Any]) -> Tuple[bool, List[str]]:
    """Validate a legal_source_register record."""
    errors = []
    if not record.get("legal_source_id"):
        errors.append("Missing legal_source_id")
    if not record.get("country_code"):
        errors.append("Missing country_code")
    if not record.get("source_type"):
        errors.append("Missing source_type")
    if not record.get("law_name_original"):
        errors.append("Missing law_name_original")
    if record.get("approval_status") not in ("not_approved", "archived_reference_only"):
        errors.append(f"Invalid approval_status: {record.get('approval_status')}")
    return len(errors) == 0, errors


def validate_legal_curated_index_record(record: Dict[str, Any]) -> Tuple[bool, List[str]]:
    """Validate a legal_curated_index record."""
    errors = []
    if not record.get("legal_index_id"):
        errors.append("Missing legal_index_id")
    if not record.get("country_code"):
        errors.append("Missing country_code")
    if not record.get("topic"):
        errors.append("Missing topic")
    if record.get("approval_status") not in ("not_approved", "reference_only"):
        errors.append(f"Invalid approval_status: {record.get('approval_status')}")
    review_status = record.get("review_status", "")
    valid_statuses = {"indexed_pending_analysis", "analysed_pending_review",
                       "confirmed", "rejected", "superseded"}
    if review_status not in valid_statuses:
        errors.append(f"Invalid review_status: {review_status}")
    return len(errors) == 0, errors


# ══════════════════════════════════════════════════════════════════
#  Write Functions
# ══════════════════════════════════════════════════════════════════


def write_legal_source_register(
    records: List[Dict[str, Any]], execute: bool = False
) -> Tuple[bool, str]:
    """Write legal_source_register records to JSONL."""
    path = legal_source_register_path()
    if not execute:
        return True, f"[DRY RUN] Would write {len(records)} records to {path}"

    ensure_empty_jsonl(path)
    with open(path, "a", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
    return True, f"Wrote {len(records)} records to {path}"


def write_legal_curated_index(
    records: List[Dict[str, Any]], execute: bool = False
) -> Tuple[bool, str]:
    """Write legal_curated_index records to JSONL."""
    path = legal_curated_index_path()
    if not execute:
        return True, f"[DRY RUN] Would write {len(records)} records to {path}"

    ensure_empty_jsonl(path)
    with open(path, "a", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
    return True, f"Wrote {len(records)} records to {path}"


# ══════════════════════════════════════════════════════════════════
#  Import Summary & Report
# ══════════════════════════════════════════════════════════════════


def generate_import_summary(
    workbook_path: str,
    batch_id: str,
    source_register_records: List[Dict],
    curated_index_records: List[Dict],
    skipped_sheets: List[str],
    unknown_sheets: List[str],
    duplicate_count: int = 0,
    missing_field_count: int = 0,
    country_profile_created: bool = False,
) -> Dict[str, Any]:
    """Generate a structured import summary."""
    return {
        "batch_id": batch_id,
        "workbook_filename": os.path.basename(workbook_path),
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "source_register_count": len(source_register_records),
        "curated_index_count": len(curated_index_records),
        "total_records": len(source_register_records) + len(curated_index_records),
        "skipped_sheets": skipped_sheets,
        "unknown_sheets": unknown_sheets,
        "duplicate_count": duplicate_count,
        "missing_field_count": missing_field_count,
        "country_profile_created": country_profile_created,
        "execution_mode": "dry_run",
        "all_not_approved": True,
        "no_legal_judgment": True,
    }


def generate_import_report(
    summary: Dict[str, Any],
    include_details: bool = True,
) -> str:
    """Generate a human-readable import report as markdown."""
    lines = []
    lines.append("# P9B Vietnam Legal Excel Import Report")
    lines.append("")
    lines.append(f"**Batch ID:** {summary.get('batch_id', 'N/A')}")
    lines.append(f"**Workbook:** {summary.get('workbook_filename', 'N/A')}")
    lines.append(f"**Timestamp:** {summary.get('timestamp', 'N/A')}")
    lines.append(f"**Execution mode:** {summary.get('execution_mode', 'N/A')}")
    lines.append("")
    lines.append("## Import Summary")
    lines.append("")
    lines.append(f"| Metric | Count |")
    lines.append(f"|--------|-------|")
    lines.append(f"| Legal Source Register records | {summary.get('source_register_count', 0)} |")
    lines.append(f"| Legal Curated Index records | {summary.get('curated_index_count', 0)} |")
    lines.append(f"| Total records | {summary.get('total_records', 0)} |")
    lines.append(f"| Skipped sheets | {len(summary.get('skipped_sheets', []))} |")
    lines.append(f"| Unknown sheets | {len(summary.get('unknown_sheets', []))} |")
    lines.append(f"| Duplicates detected | {summary.get('duplicate_count', 0)} |")
    lines.append(f"| Missing required fields | {summary.get('missing_field_count', 0)} |")
    lines.append(f"| Country profile created | {summary.get('country_profile_created', False)} |")
    lines.append("")

    if include_details:
        lines.append("## Skipped Sheets")
        for s in summary.get("skipped_sheets", []):
            lines.append(f"- {s}")
        lines.append("")
        lines.append("## Unknown Sheets")
        for s in summary.get("unknown_sheets", []):
            lines.append(f"- {s}")
        lines.append("")

    lines.append("## ⚠️ Warnings")
    lines.append("")
    lines.append("- All records are `not_approved` / pending review.")
    lines.append("- No legal judgment has been made.")
    lines.append("- No legal requirement has been confirmed as applicable.")
    lines.append("- No GSP Standard has been created, modified, or approved.")
    lines.append("- No SOP, checklist, training, or department todo has been generated.")
    lines.append("- This is a user-curated legal index — not final legal obligation data.")
    lines.append("- Legal applicability analysis and final legal review are deferred to future phases.")

    return "\n".join(lines)


def write_import_summary_json(
    summary: Dict[str, Any],
    batch_id: str,
    execute: bool = False,
) -> Tuple[bool, str]:
    """Write import summary as JSON."""
    os.makedirs(reports_dir(), exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = os.path.join(reports_dir(), f"p9b_import_summary_{ts}.json")
    if not execute:
        return True, f"[DRY RUN] Would write summary to {path}"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)
    return True, f"Wrote summary to {path}"


def write_import_report_md(
    report: str,
    batch_id: str,
    execute: bool = False,
) -> Tuple[bool, str]:
    """Write import report as markdown."""
    os.makedirs(reports_dir(), exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = os.path.join(reports_dir(), f"p9b_import_report_{ts}.md")
    if not execute:
        return True, f"[DRY RUN] Would write report to {path}"
    with open(path, "w", encoding="utf-8") as f:
        f.write(report)
    return True, f"Wrote report to {path}"


# ══════════════════════════════════════════════════════════════════
#  Country Profile
# ══════════════════════════════════════════════════════════════════


def init_vietnam_country_profile(execute: bool = False) -> Tuple[bool, str, Dict[str, Any]]:
    """Initialize the Vietnam country legal profile.

    Returns: (success, message, profile_record)
    """
    profile = {
        "country_profile_id": "CLP-VN",
        "country_code": "VN",
        "country_name": "Vietnam",
        "jurisdiction_levels": ["national", "provincial"],
        "official_languages": ["vi"],
        "working_languages": ["vi", "en", "zh"],
        "legal_categories": [
            "labour", "fire", "environment", "occupational_safety",
            "chemicals", "social_insurance", "construction", "ESG",
        ],
        "local_legal_owner": "",
        "review_required": True,
        "update_frequency": "quarterly",
        "source_priority_rules": [
            "law > decree > circular",
            "national > provincial",
        ],
        "notes": "Initialized by P9B Vietnam legal Excel import",
        "status": "active",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }

    path = country_profiles_path()
    if not execute:
        return True, f"[DRY RUN] Would write country profile CLP-VN to {path}", profile

    # Check if profile already exists
    existing = _load_jsonl(path)
    for rec in existing:
        if rec.get("country_profile_id") == "CLP-VN":
            return True, "Country profile CLP-VN already exists.", rec

    ensure_empty_jsonl(path)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(profile, ensure_ascii=False) + "\n")

    return True, f"Created country profile CLP-VN.", profile


# ══════════════════════════════════════════════════════════════════
#  Full Import Pipeline
# ══════════════════════════════════════════════════════════════════


def run_vietnam_legal_import(
    file_path: str,
    batch_id: str,
    execute: bool = False,
    max_rows: Optional[int] = None,
    sheets: Optional[List[str]] = None,
) -> Dict[str, Any]:
    """Run the full Vietnam legal Excel import pipeline.

    Args:
        file_path: Path to .xlsx file.
        batch_id: Batch identifier.
        execute: If True, writes to JSONL.
        max_rows: Optional max row limit per target.
        sheets: Optional list of sheet names to process (None = all).

    Returns:
        Dict with keys: success, message, summary, reports, warnings.
    """
    now = datetime.now(timezone.utc)
    result = {
        "success": False,
        "message": "",
        "file_path": file_path,
        "batch_id": batch_id,
        "execute": execute,
        "source_register_records": [],
        "curated_index_records": [],
        "skipped_sheets": [],
        "unknown_sheets": [],
        "warnings": [],
        "duplicate_count": 0,
        "missing_field_count": 0,
        "country_profile_created": False,
        "reports": {},
    }

    # Step 1: Inspect workbook
    try:
        import openpyxl
    except ImportError:
        result["message"] = "openpyxl not installed"
        return result

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        result["message"] = f"Failed to open workbook: {e}"
        return result

    # Step 2: Process sheets
    src_counter = len(_load_jsonl(legal_source_register_path())) + 1
    idx_counter = len(_load_jsonl(legal_curated_index_path())) + 1

    sheet_names = sheets if sheets else wb.sheetnames

    for sheet_name in wb.sheetnames:
        if sheet_names and sheet_name not in sheet_names:
            result["skipped_sheets"].append(sheet_name)
            continue

        ws = wb[sheet_name]
        target, records, cnt, warnings = _parse_sheet(ws, sheet_name, src_counter)

        if target == "legal_source_register":
            result["source_register_records"].extend(records)
            src_counter = cnt
            if max_rows and len(result["source_register_records"]) >= max_rows:
                result["source_register_records"] = result["source_register_records"][:max_rows]
        elif target == "legal_curated_index":
            result["curated_index_records"].extend(records)
            idx_counter = cnt
            if max_rows and len(result["curated_index_records"]) >= max_rows:
                result["curated_index_records"] = result["curated_index_records"][:max_rows]
        else:
            result["unknown_sheets"].append(sheet_name)

        result["warnings"].extend(warnings)

    wb.close()

    # Step 3: Validate records
    for rec in result["source_register_records"]:
        valid, errors = validate_legal_source_register_record(rec)
        if not valid:
            result["missing_field_count"] += len(errors)

    for rec in result["curated_index_records"]:
        valid, errors = validate_legal_curated_index_record(rec)
        if not valid:
            result["missing_field_count"] += len(errors)

    # Step 4: Country profile
    success, msg, profile = init_vietnam_country_profile(execute=execute)
    result["country_profile_created"] = "Created" in msg or "already exists" in msg

    # Step 5: Write records
    if execute:
        write_legal_source_register(result["source_register_records"], execute=True)
        write_legal_curated_index(result["curated_index_records"], execute=True)

    # Step 6: Generate reports
    summary = generate_import_summary(
        workbook_path=file_path,
        batch_id=batch_id,
        source_register_records=result["source_register_records"],
        curated_index_records=result["curated_index_records"],
        skipped_sheets=result["skipped_sheets"],
        unknown_sheets=result["unknown_sheets"],
        duplicate_count=result["duplicate_count"],
        missing_field_count=result["missing_field_count"],
        country_profile_created=result["country_profile_created"],
    )
    summary["execution_mode"] = "execute" if execute else "dry_run"

    report_md = generate_import_report(summary)

    # Write report files
    write_import_summary_json(summary, batch_id, execute=execute)
    write_import_report_md(report_md, batch_id, execute=execute)

    result["summary"] = summary
    result["reports"]["markdown"] = report_md
    result["success"] = True
    result["message"] = (
        f"Import {'executed' if execute else 'dry-run'} for batch {batch_id}: "
        f"{len(result['source_register_records'])} source register + "
        f"{len(result['curated_index_records'])} curated index records."
    )

    return result
